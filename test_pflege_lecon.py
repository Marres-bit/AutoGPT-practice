"""
Test de génération de leçon en allemand avec Pflege Plannung
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from medical_agent.agent import MedicalAIAgent
from medical_agent.database import MedicalDatabase
from medical_agent.config import DOCUMENT_NAME, OUTPUT_FOLDER
from datetime import datetime
from docx import Document


def main():
    print("🧪 TEST GÉNÉRATION LEÇON ALLEMANDE avec Pflege Plannung\n")
    print("=" * 60)
    
    # Créer dossier de sortie sur Bureau
    output_dir = Path.home() / "Desktop" / OUTPUT_FOLDER
    output_dir.mkdir(exist_ok=True)
    print(f"📁 Dossier de sortie: {output_dir}\n")
    
    # Initialiser agent
    agent = MedicalAIAgent()
    db = MedicalDatabase()
    
    # Générer leçon
    print("\n📚 Génération en cours...\n")
    lesson = agent.generate_lesson()
    
    if not lesson:
        print("❌ Échec génération")
        return
    
    # Sauvegarder dans DB
    db.add_lesson(
        title=lesson['title'],
        category=lesson['category'],
        difficulty=lesson['difficulty'],
        content=lesson['content']
    )
    
    # Afficher résumé
    content = lesson['content']
    print(f"\n📄 LEÇON GÉNÉRÉE:")
    print("=" * 60)
    print(f"Titre: {lesson['title']}")
    print(f"Catégorie: {lesson['category']}")
    print(f"Niveau: {lesson['difficulty']}")
    print(f"\nContenu (premiers 1000 caractères):")
    print("-" * 60)
    print(content[:1000] + "...")
    print("-" * 60)
    
    # Vérifier présence Pflege Plannung
    if "PFLEGEPLANUNG" in content.upper() or "PFLEGE" in content.upper():
        print("\n✅ PFLEGE PLANNUNG détecté dans le contenu!")
    else:
        print("\n⚠️ ATTENTION: PFLEGE PLANNUNG manquant!")
    
    # Créer document Word
    print("\n📝 Création document Word...")
    doc = Document()
    
    # En-tête
    doc.add_heading(DOCUMENT_NAME, 0)
    doc.add_heading(lesson['title'], 1)
    
    # Métadonnées
    doc.add_paragraph(f"Kategorie: {lesson['category']}")
    doc.add_paragraph(f"Schwierigkeitsgrad: {lesson['difficulty']}")
    doc.add_paragraph(f"Datum: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph()
    
    # Contenu
    doc.add_heading('Inhalt:', 2)
    for para in content.split('\n'):
        if para.strip():
            doc.add_paragraph(para)
    
    # Sauvegarder
    filename = f"{DOCUMENT_NAME}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    filepath = output_dir / filename
    doc.save(str(filepath))
    
    print(f"✅ Document sauvegardé: {filepath}")
    print(f"\n{'=' * 60}")
    print("✅ TEST RÉUSSI - Vérifiez le document sur votre Bureau!")
    print("=" * 60)


if __name__ == "__main__":
    main()
