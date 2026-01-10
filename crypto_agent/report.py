"""
Generate Word (.docx) reports summarizing market analysis and simulated trades
"""
from docx import Document
from datetime import datetime
from pathlib import Path
from typing import Dict, List


def generate_docx_report(filename: Path, market_summary: Dict, decision: Dict, trades: List[Dict], capital_state: Dict):
    doc = Document()
    doc.add_heading('SP Testnet - Rapport', level=1)

    doc.add_paragraph(f'Date: {datetime.utcnow().isoformat()} UTC')

    # Market summary
    doc.add_heading('Résumé du marché', level=2)
    doc.add_paragraph(f"Contexte général: {market_summary.get('context')}")
    doc.add_paragraph(f"Univers de trading: {market_summary.get('universe_size')}")

    doc.add_heading('Top hausses', level=3)
    for g in market_summary.get('top_gainers', [])[:10]:
        doc.add_paragraph(f"{g['symbol']}: {g['change_pct']}% (vol={g['volume']}) - {g.get('reason')}")

    doc.add_heading('Top baisses', level=3)
    for l in market_summary.get('top_losers', [])[:10]:
        doc.add_paragraph(f"{l['symbol']}: {l['change_pct']}% (vol={l['volume']}) - {l.get('reason')}")

    # Decision
    doc.add_heading('Décision SP', level=2)
    if decision:
        doc.add_paragraph(f"Action: {decision.get('action')}")
        candidate = decision.get('candidate')
        if candidate:
            doc.add_paragraph(f"Candidat choisi: {candidate.get('symbol')} - {candidate.get('change_pct')}% - raison: {candidate.get('reason')}")
    else:
        doc.add_paragraph('Aucune action prise.')

    # Trades
    doc.add_heading('Trades simulés', level=2)
    if trades:
        for t in trades:
            doc.add_paragraph(f"{t['symbol']} - entry: {t['entry_price']:.6f} qty: {t['quantity']:.6f} usd_alloc: {t['usd_allocated']:.2f} status: {t.get('status')} pnl: {t.get('pnl',0):.2f}")
    else:
        doc.add_paragraph('Aucun trade simulé')

    # Capital
    doc.add_heading('Situation du capital', level=2)
    doc.add_paragraph(f"Principal: {capital_state.get('principal'):.2f} USDT")
    doc.add_paragraph(f"Investment balance: {capital_state.get('investment_balance'):.2f} USDT")
    doc.add_paragraph(f"Cumulative profits: {capital_state.get('cumulative_profits'):.2f} USDT")

    # Save
    filename.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(filename))
    return filename
