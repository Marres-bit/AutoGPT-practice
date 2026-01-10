"""
Module de Génération de Rapports Word Automatiques
Crée/met à jour les rapports crypto sur le Bureau
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from pathlib import Path
from datetime import datetime
from typing import List, Dict


class WordReporter:
    """Génère et met à jour les rapports Word des analyses crypto"""
    
    def __init__(self, desktop_path=None):
        """
        Initialise le générateur de rapports
        
        Args:
            desktop_path: Chemin du Bureau (défaut: Bureau de l'utilisateur)
        """
        if desktop_path is None:
            self.desktop_path = Path.home() / "Desktop"
        else:
            self.desktop_path = Path(desktop_path)
        
        self.desktop_path.mkdir(parents=True, exist_ok=True)
    
    def get_report_filename(self, date=None):
        """Génère le nom du fichier rapport"""
        if date is None:
            date = datetime.now()
        
        return self.desktop_path / f"Crypto_Compte_rendu_{date.strftime('%Y-%m-%d')}.docx"
    
    def create_or_update_report(self, analysis_data: Dict) -> str:
        """
        Crée ou met à jour le rapport Word du jour
        
        Args:
            analysis_data: Dictionnaire avec les analyses (from crypto_analyzer)
        
        Returns:
            Chemin du fichier créé/mis à jour
        """
        report_path = self.get_report_filename()
        
        # Charger le document existant ou en créer un nouveau
        if report_path.exists():
            print(f"📄 Mise à jour du rapport: {report_path.name}")
            doc = Document(report_path)
        else:
            print(f"📄 Création du rapport: {report_path.name}")
            doc = self._create_new_document()
        
        # Ajouter les analyses du jour
        self._add_analysis_section(doc, analysis_data)
        
        # Sauvegarder
        doc.save(report_path)
        print(f"✅ Rapport sauvegardé: {report_path}")
        
        return str(report_path)
    
    def _create_new_document(self) -> Document:
        """Crée un nouveau document Word avec formatage"""
        doc = Document()
        
        # Titre principal
        title = doc.add_heading("📊 Rapports d\'Analyses Crypto Autonomes", 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Sous-titre
        subtitle = doc.add_paragraph(
            "Analyseur de crypto-monnaies ultra autonome | Mises à jour quotidiennes"
        )
        subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        subtitle_format = subtitle.runs[0]
        subtitle_format.italic = True
        subtitle_format.font.size = Pt(10)
        
        # Ligne de séparation
        doc.add_paragraph("_" * 80)
        
        return doc
    
    def _add_analysis_section(self, doc: Document, analysis_data: Dict):
        """Ajoute une section d'analyse au document"""
        
        # En-tête avec date/heure
        timestamp = analysis_data.get("timestamp", datetime.now().isoformat())
        date_str = datetime.fromisoformat(timestamp).strftime("%d/%m/%Y à %H:%M:%S")
        
        header = doc.add_heading(f"📅 Analyse du {date_str}", level=1)
        header_run = header.runs[0]
        header_run.font.color.rgb = RGBColor(0, 102, 204)
        
        # Résumé rapide
        cryptos = analysis_data.get("cryptos", [])
        if cryptos:
            gainers = [c for c in cryptos if "+" in c.get("change_24h", "")]
            losers = [c for c in cryptos if "-" in c.get("change_24h", "")]
            
            summary = doc.add_paragraph()
            summary.add_run(f"📈 Cryptos en hausse: {len(gainers)} | ")
            summary.add_run(f"📉 Cryptos en baisse: {len(losers)}\n")
            summary.add_run(f"⏱️ Total: {len(cryptos)} cryptos analysées")
        
        doc.add_paragraph()  # Espace
        
        # Tableau avec les analyses
        if cryptos:
            table = doc.add_table(rows=1, cols=7)
            table.style = 'Light Grid Accent 1'
            
            # En-tête du tableau
            header_cells = table.rows[0].cells
            headers = ["Crypto", "Prix", "Variation 24h", "Rang", "Liquidité", "Volume/Cap", "Analyse"]
            for i, header_text in enumerate(headers):
                cell = header_cells[i]
                cell.text = header_text
                # Style en-tête
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    paragraph.paragraph_format.space_before = Pt(6)
                    paragraph.paragraph_format.space_after = Pt(6)
            
            # Lignes avec données
            for crypto in cryptos:
                row_cells = table.add_row().cells
                
                row_cells[0].text = f"{crypto['symbol']}\n{crypto['name']}"
                row_cells[1].text = crypto.get('price', 'N/A')
                row_cells[2].text = crypto.get('change_24h', 'N/A')
                row_cells[3].text = str(crypto.get('rank', 'N/A'))
                row_cells[4].text = crypto.get('liquidity', 'N/A')
                row_cells[5].text = crypto.get('volume_ratio', 'N/A')
                
                analysis_text = f"{crypto.get('strength', '')}\n{crypto.get('analysis', '')}"
                row_cells[6].text = analysis_text
                
                # Colorer les cellules de variation
                change_text = crypto.get('change_24h', '')
                if '+' in change_text:
                    row_cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 153, 0)
                elif '-' in change_text:
                    row_cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(204, 0, 0)
        
        # Pied de page
        doc.add_paragraph()
        footer = doc.add_paragraph(
            f"Rapport généré automatiquement le {date_str} | Données de CoinGecko API"
        )
        footer_run = footer.runs[0]
        footer_run.font.size = Pt(8)
        footer_run.font.italic = True
        footer_run.font.color.rgb = RGBColor(128, 128, 128)
        
        # Séparateur
        doc.add_paragraph("_" * 80)
        doc.add_paragraph()  # Espace pour prochaine analyse


def main():
    """Test du module"""
    reporter = WordReporter()
    
    # Données de test
    test_data = {
        "timestamp": datetime.now().isoformat(),
        "cryptos": [
            {
                "name": "Bitcoin",
                "symbol": "BTC",
                "price": "$45,000.00",
                "change_24h": "+5.50%",
                "rank": 1,
                "liquidity": "💧 Très bonne",
                "volume_ratio": "45.2%",
                "strength": "📈 FORT",
                "analysis": "Hausse positive avec volume soutenu"
            },
            {
                "name": "Ethereum",
                "symbol": "ETH",
                "price": "$2,500.00",
                "change_24h": "-3.20%",
                "rank": 2,
                "liquidity": "💧 Très bonne",
                "volume_ratio": "38.1%",
                "strength": "📉 FORT",
                "analysis": "Baisse temporaire, zone de support importante"
            }
        ]
    }
    
    path = reporter.create_or_update_report(test_data)
    print(f"✅ Rapport créé: {path}")


if __name__ == "__main__":
    main()
