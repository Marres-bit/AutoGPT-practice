"""
Script pour exporter toutes les leçons médicales en fichiers Word
"""
import sqlite3
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from datetime import datetime

# Chemin du bureau
DESKTOP = Path.home() / "Desktop"
OUTPUT_FOLDER = DESKTOP / "Leçons_Médicales"

def export_lessons_to_word():
    """Exporte toutes les leçons en fichiers Word sur le bureau"""
    
    # Créer le dossier de sortie
    OUTPUT_FOLDER.mkdir(exist_ok=True)
    print(f"📁 Dossier créé: {OUTPUT_FOLDER}")
    
    # Connexion à la base de données
    conn = sqlite3.connect('medical_lessons.db')
    cursor = conn.cursor()
    
    # Récupérer toutes les leçons
    cursor.execute('''
        SELECT id, title, category, difficulty, duration, content, generated_at, views, rating
        FROM lessons
        ORDER BY id
    ''')
    
    lessons = cursor.fetchall()
    
    if not lessons:
        print("❌ Aucune leçon trouvée dans la base de données")
        return
    
    print(f"\n📚 Export de {len(lessons)} leçons en cours...\n")
    
    for lesson in lessons:
        lesson_id, title, category, difficulty, duration, content, generated_at, views, rating = lesson
        
        # Créer un document Word
        doc = Document()
        
        # === EN-TÊTE ===
        header = doc.add_heading('🧠 MediGenius AI', level=0)
        header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Titre de la leçon
        title_para = doc.add_heading(f'Leçon #{lesson_id}: {category}', level=1)
        title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        doc.add_paragraph()
        
        # === MÉTADONNÉES ===
        metadata_table = doc.add_table(rows=5, cols=2)
        metadata_table.style = 'Light Grid Accent 1'
        
        # Remplir le tableau de métadonnées
        metadata = [
            ('📚 Catégorie', category or 'N/A'),
            ('🎯 Spécialité', title or 'N/A'),
            ('📊 Difficulté', difficulty or 'N/A'),
            ('⏱️ Durée', duration or 'N/A'),
            ('📅 Généré le', generated_at or 'N/A')
        ]
        
        for i, (label, value) in enumerate(metadata):
            metadata_table.rows[i].cells[0].text = label
            metadata_table.rows[i].cells[1].text = str(value)
            # Mettre en gras la première colonne
            metadata_table.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        
        doc.add_paragraph()
        
        # === STATISTIQUES ===
        stats_para = doc.add_paragraph()
        stats_para.add_run(f'👁️ Vues: {views or 0}').bold = True
        stats_para.add_run(f'  |  ')
        stats_para.add_run(f'⭐ Note: {rating or 0.0}/5').bold = True
        stats_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        doc.add_paragraph()
        doc.add_paragraph('─' * 80)
        doc.add_paragraph()
        
        # === CONTENU ===
        content_heading = doc.add_heading('📖 Contenu de la Leçon', level=2)
        
        if content and len(content) > 10:
            # Si le contenu est long, on le divise en paragraphes
            paragraphs = content.split('\n\n')
            for para in paragraphs:
                if para.strip():
                    p = doc.add_paragraph(para.strip())
                    p.paragraph_format.line_spacing = 1.5
                    p.paragraph_format.space_after = Pt(12)
        else:
            doc.add_paragraph(content or 'Contenu non disponible')
        
        doc.add_paragraph()
        doc.add_paragraph('─' * 80)
        
        # === FOOTER ===
        footer = doc.add_paragraph()
        footer.add_run(f'\n📅 Document généré le {datetime.now().strftime("%d/%m/%Y à %H:%M")}').italic = True
        footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Nom du fichier sécurisé
        safe_category = "".join(c for c in category if c.isalnum() or c in (' ', '-', '_')).strip()
        safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '-', '_')).strip()
        filename = f"Leçon_{lesson_id:02d}_{safe_category}_{safe_title}.docx"
        filepath = OUTPUT_FOLDER / filename
        
        # Sauvegarder le document
        doc.save(filepath)
        print(f"✅ {filename}")
    
    conn.close()
    
    print(f"\n🎉 Export terminé!")
    print(f"📂 Emplacement: {OUTPUT_FOLDER}")
    print(f"📄 {len(lessons)} fichiers Word créés")

if __name__ == "__main__":
    print("🧠 MediGenius AI - Export des Leçons")
    print("=" * 60)
    export_lessons_to_word()
